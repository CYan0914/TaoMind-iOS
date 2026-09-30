# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Arial'
style.font.size = Pt(11)
style.paragraph_format.line_spacing = 1.35
style.paragraph_format.space_after = Pt(4)

def div():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run('━' * 50)
    r.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    r.font.size = Pt(7)

def h2(text):
    h = doc.add_heading(text, level=2)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

def h3(text):
    h = doc.add_heading(text, level=3)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x55)

def b(p, text, size=11):
    r = p.add_run(text)
    r.font.bold = True
    r.font.size = Pt(size)
    return r

def rn(p, text, size=10.5, color=None, italic=False, bold=False):
    r = p.add_run(text)
    r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if italic: r.font.italic = True
    if bold: r.font.bold = True
    return r

def bullet(text):
    doc.add_paragraph(text, style='List Bullet')

def quote(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.right_indent = Cm(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0x44, 0x44, 0x66)
    return p

# ============ TITLE ============
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run(u'中金（上海）· 量化策略工程师（J19485）\n面试准备稿')
r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
rn(sub, u'信息技术部 · 社会招聘 · 内推渠道', size=12, italic=True, color=RGBColor(0x66, 0x66, 0x66))

doc.add_paragraph()
div()

# ============ 0 ============
h2(u'〇、先读懂这位面试官')

p = doc.add_paragraph()
rn(p, u'你的同学提供的关键背景——这位领导（技术负责人）的人格画像：', size=11, bold=True)

profile = [
    u'2017-18年曾经历职业生涯低谷：因与上级领导理念不合被边缘化——他懂得"理念冲突"和"职业低谷"意味着什么',
    u'当时孩子读小学，他每周上海↔北京往返，坚持了一年——他是顾家、重家庭的人，会理解"为了孩子做决定"的分量',
    u'能从低谷走出来并成为中金的技术负责人——他欣赏有韧性、能自我驱动、能把故事讲成成长的人',
]
for pt in profile:
    bullet(pt)

p = doc.add_paragraph()
rn(p, u'这意味着什么：', size=11, bold=True)
rn(p, u'这位面试官会用"过来人"的视角看你。他不在乎你粉饰的完美履历，而在乎你有没有把低谷讲成成长的能力。所以你的"宁波→上海"和"博士肄业"这两件事，恰恰是他最容易共情的点——讲好了是加分项，不是减分项。', size=11, color=RGBColor(0xCC, 0x33, 0x00))

doc.add_paragraph()
div()

# ============ 1 ============
h2(u'一、核心问题①：为什么选择从宁波来上海？')

p = doc.add_paragraph()
rn(p, u'这个问题几乎100%会问。你的同学建议用"为了孩子"这个角度——非常对路，因为这位面试官自己就是"为了家庭往返两地一年"的人。', size=11)

p = doc.add_paragraph()
rn(p, u'但要讲得高级，得是"家庭动因 + 职业动因"两条腿，缺一条都单薄。', size=11, bold=True)

h3(u'逐字稿（建议30-40秒）：')

quote(u'"说实话，这个决定我考虑了挺长时间。第一个原因是家庭——我孩子刚半岁，我和太太商量了很久，觉得宁波虽然生活安逸，但孩子将来读书、户口这些，上海的综合资源确实更好。我现在的核心诉求就是给孩子一个更稳定、更好的成长环境，这件事在我们家是共识。第二个原因是职业——量化这条赛道，宁波的机会和上海完全不在一个量级。中金的平台、团队、技术积累，对我来说是天花板级别的机会。我判断这个行业的趋势就是往头部平台集中，所以我愿意带着全家搬过来。"')

p = doc.add_paragraph()
rn(p, u'这段话的三个设计：', size=10.5, bold=True)
design = [
    u'"考虑了挺长时间" → 显示你深思熟虑，不是冲动搬迁',
    u'"我们家的共识" → 家庭支持，没有后顾之忧，能长期稳定',
    u'"天花板级别的机会" → 不是单纯投奔家庭，而是职业判断，两条腿都站得住',
]
for d in design:
    bullet(d)

p = doc.add_paragraph()
rn(p, u'⚠ 对方可能追问：', size=10.5, bold=True, color=RGBColor(0xCC, 0x66, 0x00))
follow = [
    u'"那你太太工作怎么办？" → "她现在在宁波，我们计划在上海安顿稳定后，看她有没有合适的远程或本地机会。孩子还小，这两年我先把这边站稳。"',
    u'"会不会干两年又回宁波了？" → "不会。我这次是全家一起搬，不是一个人来试水。房子和孩子读书都会落实，是奔着长期来的。"',
]
for f in follow:
    bullet(f)

doc.add_paragraph()
div()

# ============ 2 ============
h2(u'二、核心问题②：博士肄业这件事，怎么讲成故事？')

p = doc.add_paragraph()
rn(p, u'你的简历写的是"因国内职业发展机会提前离校"——这是事实，但太干。同学建议编一个有情怀的故事，方向是对的，但核心原则是：', size=11, bold=True)
rn(p, u'故事要真、要有取舍、要有成长，不能编造你博士期间没发生的事。', size=11, color=RGBColor(0xCC, 0x33, 0x00))

h3(u'讲故事的三个层次：')

layers = [
    (u'第一层：为什么离校（要讲成"选择"而不是"放弃"）',
     u'"当时在华威读博，研究方向偏学术。但我在写论文的时候，越来越清楚地感觉到——我喜欢的是把金融模型放到真实市场里去验证，而不是在理论上推导。国内正好有一个做买方研究的实际机会，可以真金白银地做投资决策。我衡量了很久，决定抓住这个机会。说实话，放弃博士学位是当时我做过最纠结的决定之一。"'),
    (u'第二层：代价与成长（诚实承认但不自怜）',
     u'"代价是有的——博士没读完，在很多人看来是个遗憾。但回头看，这几年我把时间花在了真实的市场里，从买方研究做到基金管理，管过钱、跑过矿山、做过量化。如果当时留在学校，可能到现在还在发论文，但我不确定那是我真正想要的。"'),
    (u'第三层：落地到现在这个岗位（闭环）',
     u'"现在回头看，当时的判断是对的。我现在更想做的是把投资经验转化成工程化的量化工具，而中金这个岗位正好是让技术真正驱动投资决策的地方。所以我觉得这个选择一路把我带到了这里。"'),
]
for title_text, content in layers:
    p = doc.add_paragraph()
    b(p, title_text, 10.5)
    quote(content)

p = doc.add_paragraph()
rn(p, u'⚠ 注意事项：', size=10.5, bold=True, color=RGBColor(0xCC, 0x66, 0x00))
notes = [
    u'绝对不要抱怨学术体系、导师或华威——面试官会立刻警觉',
    u'"放弃博士学位"这句话要说成"我选择了一条不一样的路"，不要流露出懊悔或自我否定',
    u'如果面试官是海归，他会懂PhD drop的真实成本，你的"纠结"他会感同身受——这反而是拉近距离的机会',
    u'千万不要提到任何"家庭经济原因""读不下去"之类的暗示，会把故事讲low',
]
for n in notes:
    bullet(n)

doc.add_paragraph()
div()

# ============ 3 ============
h2(u'三、核心问题③：为什么从基金经理转量化策略工程师？')

p = doc.add_paragraph()
rn(p, u'这位面试官是技术出身，你要让他相信：你不是因为做投资做不下去才转技术，而是"看得更远"。你的同学给的切入点很好——股票市场量化横行。', size=11)

h3(u'逐字稿：')

quote(u'"我做股票投资这几年，最深的体会是市场结构变了。现在A股的主观多头，面对的对手盘大量是程序化、量化的资金，信息效率被大幅拉高，主观交易的超额收益空间在收窄——这不是说主观不行，而是说未来的投资，越来越依赖工程化的工具和数据基础设施。我自己这几年一直在做量化系统，从因子构建到回测平台都是自己搭的。我就想，与其在投资端跟量化抢钱，不如去做造工具、造平台的那一端。中金自研算法执行平台和算法总线，就是把技术和投资结合起来的地方，这正好是我最想做的事。"')

p = doc.add_paragraph()
rn(p, u'这段话的巧妙之处：', size=10.5, bold=True)
design2 = [
    u'承认市场变化（量化横行）是客观判断，不是能力不足——你先主动说出来，对方就没法拿这个攻击你',
    u'你的量化背景（自建回测平台）是桥梁——证明你已经是"懂技术的投资人"，转工程不是从零开始',
    u'"造工具的人"这个定位——把转岗讲成产业趋势判断，格局一下就上来了',
]
for d in design2:
    bullet(d)

p = doc.add_paragraph()
rn(p, u'⚠ 对方可能追问：', size=10.5, bold=True, color=RGBColor(0xCC, 0x66, 0x00))
follow2 = [
    u'"你的Java/C++水平怎么样？" → 诚实："我的主力语言是Python，金融建模、回测都用它。C++读过大学课程（85分），有面向对象基础，但生产级经验确实没有。Java我最近在系统学习。我学习能力很强，如果中金给我这个机会，我可以在最短时间上手。"（不要吹，技术面试官一眼看穿）',
    u'"你会不会做一段时间又回去做投资？" → "不会。我投资的经验恰恰是我做这个岗位的差异化优势——我懂量化策略的真实需求，知道平台该支持什么。这是我这个岗位区别于纯技术出身候选人的地方。"',
]
for f in follow2:
    bullet(f)

doc.add_paragraph()
div()

# ============ 4 ============
h2(u'四、针对JD的专业准备')

h3(u'4.1 拆解这4条工作职责——你想清楚了，面试就赢了')

jd = [
    (u'职责1：算法执行平台的开发、运维、迭代',
     u'这是交易下单/执行系统。你不需要做过，但要能说出：执行平台要解决的是"怎么把策略信号变成最优成交"——包括拆单、时机、冲击成本。你可以结合自己回测的经验说：回测和实盘最大的差别是滑点和冲击成本，执行平台就是处理这个的。'),
    (u'职责2：算法总线的开发、维护、性能调优',
     u'这是把策略和外部数据/系统连接起来的中枢。性能调优意味着高吞吐、低延时。你虽然没做过，但你在金融建模里对性能敏感（大规模回测、因子计算）——这是可以迁移的点。'),
    (u'职责3：量化策略配套支撑体系（数据、回测框架）',
     u'这块你是真正的内行！你自己搭过完整回测系统。这一条你要重点讲，把你自建回测系统的经验展开——数据清洗、因子库、回测引擎、绩效归因。这是你比纯技术候选人强的地方。'),
    (u'职责4：策略平台建设与优化',
     u'从0到1的平台思维。你独立搭过量化系统、开发过AI数字人系统——你理解"平台"意味着什么：可扩展、可复用、面向多用户。'),
]
for title_text, content in jd:
    p = doc.add_paragraph()
    b(p, title_text, 10.5)
    rn(p, content, 10.5)

doc.add_paragraph()

h3(u'4.2 必须提前做的功课（今晚+明天）')

prep = [
    u'了解中金的自研算法平台：搜"中金 算法执行 平台""中金 量化 自研系统"，看看公开资料里中金信息技术部在做什么',
    u'搞清楚DolphinDB是什么：高并发分布式时序数据库，量化领域常用。你简历写了"有基础"——但你要能说出它和MySQL/Pandas的区别，以及为什么量化场景用它（列存储+内置函数+流式计算）',
    u'复习Java/C++基础：面试官很可能出一道简单的编程题（数组/字符串/排序）。你Python熟，但面试可能要求Java或C++写',
    u'准备一个"你搭过的量化系统"的架构图：回测系统有哪些模块、数据怎么流动——画得出图就说明你真懂',
    u'准备一个低延时/高并发的概念理解：消息队列（如Kafka）、零拷贝、内存池、无锁队列——能说出两三个关键词并解释清楚即可',
]
for p_text in prep:
    bullet(p_text)

doc.add_paragraph()

h3(u'4.3 技术面高频问题速答')

tech = [
    (u'Q：回测框架最核心的几个模块？',
     u'"数据层（行情、财务、因子）、信号层（因子计算→选股）、组合层（权重、约束）、交易层（撮合、滑点、手续费）、绩效层（收益曲线、回撤、IC、归因）。我搭的时候最花精力的是数据一致性和交易模拟的真实性——这两块直接决定回测结果能不能参考。"'),
    (u'Q：Python和C++在量化里分别怎么用？',
     u'"Python做研究和原型——开发快、生态好，Pandas/NumPy/Scikit-learn都是标配。C++/Java做生产环境——执行速度、低延时要求高。我自己现在是Python为主，但理解两者的边界：策略研究重迭代，生产执行重性能。"'),
    (u'Q：你知道什么是算法总线吗？',
     u'"按我的理解，算法总线是把各类算法策略（TWAP/VWAP等执行算法、以及各类交易策略）和外部系统（行情、柜台、风控、数据）统一接入的中枢，类似消息总线的思想。它的关键是低延时和可靠性——交易场景对这两点的要求是毫秒级的。"'),
    (u'Q：多因子选股里，你怎么处理过拟合？',
     u'"三个手段：一是样本外验证，留出时间段不参与调参；二是因子间相关性控制和降维，避免冗余；三是参数敏感性分析——一个参数小幅变动结果剧烈波动，说明模型不稳。我搭的策略都会做这三步。"'),
]
for q, a in tech:
    p = doc.add_paragraph()
    b(p, q, 10.5)
    rn(p, a, 10.5)
    doc.add_paragraph()

div()

# ============ 5 ============
h2(u'五、把三个故事串起来——完整开场自我介绍（60-90秒）')

quote(u'"我简单介绍一下我自己。我的经历可能比较特别——华威金融数学硕士毕业，后来继续读博，但读了一年，我选择了一条不一样的路：回国做真正的投资。当时有人觉得可惜，但我觉得把模型放进真实市场验证，比在理论上推导更让我兴奋。这几年我做买方研究，从有色金属的深度研究，到后来独立管理2000万的股票基金，中间一直没放下编程——我的选股系统、回测平台都是自己用Python搭的。现在市场结构变了，量化资金占比越来越高，我越来越清楚自己想做什么：不是去跟量化抢钱，而是去做造工具、造平台的人，用工程能力支撑投资决策。中金自研的算法执行平台和策略支撑体系，就是这个方向上最好的机会之一。这是我特别想来的原因。"')

p = doc.add_paragraph()
rn(p, u'这段开场把三件事全部埋进去了：', size=10.5, bold=True)
embed = [
    u'博士离校 → 讲成"选择"，有情怀，不卑不亢',
    u'宁波→上海 → 没细讲，留给对方问，你再用"孩子+职业"回答',
    u'转量化工程 → 讲成产业趋势判断，顺势引出你对这个岗位的理解',
]
for e in embed:
    bullet(e)

doc.add_paragraph()
div()

# ============ 6 ============
h2(u'六、反问环节——要问得有水平')

p = doc.add_paragraph()
rn(p, u'技术负责人面试，你的反问也要偏技术/平台，不要问福利待遇。', size=11)

counter = [
    u'"中金自研算法执行平台目前处于什么阶段？是从零搭建还是已经有基座、在做迭代？" → 判断是开荒还是加入成熟团队，同时显示你做过功课',
    u'"团队目前多少人？算法总线和执行平台的开发是分开的团队还是一个团队？" → 了解组织结构和你的位置',
    u'"咱们平台现在主要服务内部自营还是也会对外输出？" → 展示你对量化技术商业化的思考',
    u'"您在这个岗位上，最希望新同事在哪个方面补上团队的短板？" → 直接摸需求，展示主动性',
]
for c in counter:
    bullet(c)

doc.add_paragraph()
div()

# ============ 7 ============
h2(u'七、面试当天提醒')

reminders = [
    u'技术面试官更在意"你懂不懂自己在说什么"，比在意你简历写了什么更重要——三个故事要讲到自然，不能像背稿',
    u'他问你宁波→上海时，先讲家庭（孩子），再讲职业，顺序别反——先家庭能迅速建立共情，再职业展示理性',
    u'博士肄业主动带出来，不要等对方发现简历里有个刺——你主动讲成"选择"，就化解了',
    u'回答不了的技术问题，说"这块我确实没做过，但我理解它的核心是XX，我学习能力强"——技术人不讨厌不会，讨厌装懂',
    u'他讲过自己上海北京往返一年的经历，如果你能自然提到"我很理解这种为了家庭两地奔波的选择"——这是无声的共情，他会记得你',
]
for rem in reminders:
    bullet(rem)

doc.add_paragraph()
div()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
rn(p, u'祝面试顺利。你手上的牌比你想象的好——懂投资的技术候选人，在中金的信息技术部是稀缺的。', size=11, italic=True)

# Save
desktop = os.path.expanduser('~/Desktop')
path = os.path.join(desktop, u'中金_量化策略工程师_面试准备稿.docx')
doc.save(path)
print(f'Done: {path}')
